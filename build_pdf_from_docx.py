from pathlib import Path
from xml.sax.saxutils import escape
from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

source = Path('output/TaskFlow_Cloud_Report_without_S3.docx')
target = Path('output/TaskFlow_Cloud_Report_without_S3.pdf')
docx = Document(source)
styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name='BodyClean', parent=styles['BodyText'], fontName='Helvetica', fontSize=9.3, leading=12, spaceAfter=5))
styles.add(ParagraphStyle(name='H1Clean', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=14, leading=17, spaceBefore=12, spaceAfter=6, keepWithNext=True))
styles.add(ParagraphStyle(name='H2Clean', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=11.2, leading=14, spaceBefore=8, spaceAfter=4, keepWithNext=True))
styles.add(ParagraphStyle(name='TitleClean', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=19, leading=23, alignment=TA_CENTER, spaceAfter=8))
styles.add(ParagraphStyle(name='CenterClean', parent=styles['BodyText'], fontName='Helvetica', fontSize=10, leading=13, alignment=TA_CENTER, spaceAfter=3))
styles.add(ParagraphStyle(name='BulletClean', parent=styles['BodyText'], fontName='Helvetica', fontSize=9.3, leading=12, leftIndent=15, firstLineIndent=-9, bulletIndent=5, spaceAfter=3))

story=[]
para_by_element = {p._p: p for p in docx.paragraphs}
table_by_element = {t._tbl: t for t in docx.tables}

def append_paragraph(p):
    text = p.text.strip()
    if not text:
        return
    n = p.style.name
    e = escape(text).replace('\n','<br/>')
    if n == 'Title': story.append(Paragraph(e, styles['TitleClean']))
    elif n == 'Heading 1': story.append(Paragraph(e, styles['H1Clean']))
    elif n == 'Heading 2': story.append(Paragraph(e, styles['H2Clean']))
    elif 'List Bullet' in n: story.append(Paragraph(e, styles['BulletClean'], bulletText='•'))
    elif p.alignment == 1: story.append(Paragraph(e, styles['CenterClean']))
    else: story.append(Paragraph(e, styles['BodyClean']))

def append_table(tbl):
    rows=[]
    for row in tbl.rows:
        rows.append([Paragraph(escape(c.text).replace('\n','<br/>'), styles['BodyClean']) for c in row.cells])
    col_width=(A4[0]-1.3*inch)/len(rows[0])
    table=Table(rows, colWidths=[col_width]*len(rows[0]), repeatRows=1, hAlign='LEFT')
    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),0.35,colors.grey),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#D9EAF7')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3)]))
    story.extend([Spacer(1, 5), table, Spacer(1, 8)])

for element in docx.element.body.iterchildren():
    if element.tag.endswith('}p') and element in para_by_element:
        append_paragraph(para_by_element[element])
    elif element.tag.endswith('}tbl') and element in table_by_element:
        append_table(table_by_element[element])

def footer(canvas, document):
    canvas.saveState(); canvas.setFont('Helvetica', 8); canvas.drawCentredString(A4[0]/2, 0.4*inch, f'TaskFlow Cloud Report | {document.page}'); canvas.restoreState()

pdf = SimpleDocTemplate(str(target), pagesize=A4, rightMargin=.65*inch, leftMargin=.65*inch, topMargin=.65*inch, bottomMargin=.6*inch)
pdf.build(story, onFirstPage=footer, onLaterPages=footer)
