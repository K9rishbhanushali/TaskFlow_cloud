from docx import Document
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from pathlib import Path

SRC = Path('output/docx/TaskFlow_Cloud_Report.docx')
OUT = Path('output/pdf/TaskFlow_Cloud_Report.pdf')
OUT.parent.mkdir(parents=True, exist_ok=True)
docx = Document(SRC)
styles = getSampleStyleSheet()
body = ParagraphStyle('body', parent=styles['BodyText'], fontName='Times-Roman', fontSize=10, leading=13, alignment=TA_JUSTIFY, spaceAfter=6)
head1 = ParagraphStyle('h1', parent=body, fontName='Times-Bold', fontSize=15, leading=19, spaceBefore=10, spaceAfter=8, alignment=TA_CENTER)
head2 = ParagraphStyle('h2', parent=body, fontName='Times-Bold', fontSize=12, leading=15, spaceBefore=8, spaceAfter=5)
caption = ParagraphStyle('caption', parent=body, fontName='Times-Italic', fontSize=9, leading=11, alignment=TA_CENTER, spaceAfter=8)
cover = ParagraphStyle('cover', parent=body, fontName='Times-Bold', fontSize=18, leading=22, alignment=TA_CENTER, textColor=colors.HexColor('#000080'))

def has_break(p):
    return bool(p._p.xpath('.//w:br[@w:type="page"]'))
def as_para(text, style=body):
    return Paragraph((text or '').replace('&','&amp;').replace('<','&lt;').replace('>','&gt;'), style)
def table_flow(tbl):
    data=[]
    for ri,row in enumerate(tbl.rows):
        vals=[]
        for c in row.cells:
            text='\n'.join(p.text for p in c.paragraphs).strip()
            vals.append(as_para(text, body if ri else ParagraphStyle('th', parent=body, fontName='Times-Bold', alignment=TA_CENTER)))
        data.append(vals)
    # screenshot marker table
    raw=' '.join(' '.join(p.text for p in c.paragraphs) for r in tbl.rows for c in r.cells)
    if '{screeshots}' in raw:
        t=Table([[as_para('{screeshots}', ParagraphStyle('ss', parent=body, alignment=TA_CENTER, fontSize=13))]], colWidths=[6.5*inch], rowHeights=[1.2*inch])
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#F2F2F2')),('GRID',(0,0),(-1,-1),.5,colors.grey),('VALIGN',(0,0),(-1,-1),'MIDDLE')]))
        return t
    n=len(data[0]) if data else 1
    widths = [6.5*inch/n]*n
    t=Table(data, colWidths=widths, repeatRows=1, hAlign='CENTER')
    commands=[('GRID',(0,0),(-1,-1),.35,colors.HexColor('#B7B7B7')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]
    if len(data)>1: commands += [('BACKGROUND',(0,0),(-1,0),colors.HexColor('#D9E2F3'))]
    t.setStyle(TableStyle(commands)); return t

story=[]
# Build a clean, single-page cover independently of the editable DOCX structure.
story.extend([Spacer(1,.45*inch), as_para("Mahavir Education Trust's", ParagraphStyle('ct1', parent=cover, fontSize=10, leading=14)), as_para('Shah & Anchor Kutchhi Engineering College', cover), as_para('An Autonomous Institute Affiliated to University of Mumbai', ParagraphStyle('ct2', parent=cover, fontSize=9, leading=12)), Spacer(1,.3*inch), as_para('UG Program in Information Technology', ParagraphStyle('ct3', parent=cover, fontSize=15, leading=18)), Spacer(1,.25*inch)])
cover_data=[[as_para('Report', ParagraphStyle('cvh',parent=body,fontName='Times-Bold',alignment=TA_CENTER)),''],[as_para('Course Name:',body),as_para('CLOUD COMPUTING',body)],[as_para('Title :',body),as_para('TaskFlow Cloud: A Cloud Based Student Task and Study Management System',body)],[as_para('Group Members Details',body),''],[as_para('Roll Number',body),as_para('Name of Student',body)],[as_para('124BTIT1039',body),as_para('KRISH BHANUSHALI',body)],[as_para('124BTIT1070',body),as_para('HET FINAVIA',body)]]
ct=Table(cover_data,colWidths=[3.05*inch,3.45*inch],hAlign='CENTER')
ct.setStyle(TableStyle([('SPAN',(0,0),(1,0)),('SPAN',(0,3),(1,3)),('GRID',(0,0),(-1,-1),.5,colors.black),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#D9E2F3')),('BACKGROUND',(0,3),(-1,3),colors.HexColor('#F2F2F2')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
story.extend([ct,PageBreak()])
past_cover=False
for child in docx.element.body.iterchildren():
    if child.tag == qn('w:p'):
        pnode = next((p for p in docx.paragraphs if p._p is child), None)
        if pnode is None: continue
        text=pnode.text.strip()
        if not past_cover:
            if has_break(pnode): past_cover=True
            continue
        if has_break(pnode):
            if text: story.append(as_para(text, body))
            story.append(PageBreak()); continue
        if not text: continue
        name=pnode.style.name
        if name == 'Heading 1': story.append(as_para(text, head1))
        elif name == 'Heading 2': story.append(as_para(text, head2))
        elif name.startswith('List Bullet'): story.append(as_para('-  '+text, body))
        elif name.startswith('List Number'): story.append(as_para('-  '+text, body))
        elif text.startswith('Figure '): story.append(as_para(text, caption))
        elif 'Shah & Anchor Kutchhi Engineering College' in text: story.append(as_para(text.replace('\n','<br/>'), cover))
        else: story.append(as_para(text, body))
    elif child.tag == qn('w:tbl'):
        if not past_cover: continue
        tbl = next((t for t in docx.tables if t._tbl is child), None)
        if tbl: story.extend([table_flow(tbl), Spacer(1,8)])

def footer(canvas, d):
    canvas.saveState(); canvas.setFont('Times-Roman',8); canvas.setFillColor(colors.grey)
    canvas.drawCentredString(letter[0]/2, .34*inch, f'TaskFlow Cloud Report | {d.page}')
    canvas.restoreState()

pdf = SimpleDocTemplate(str(OUT), pagesize=letter, leftMargin=.72*inch, rightMargin=.72*inch, topMargin=.65*inch, bottomMargin=.62*inch)
pdf.build(story, onFirstPage=footer, onLaterPages=footer)
print(OUT)
